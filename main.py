import asyncio
import os
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from aiogram import Bot, Dispatcher, F, types
from aiogram.filters import CommandStart
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.utils.keyboard import InlineKeyboardBuilder, ReplyKeyboardBuilder

TOKEN = "8531045283:AAEXf-LTfkux_HdFa5UhTPohAbK4l48nS2I"
ADMIN_ID = 5635839198  # آیدی عددی تلگرام شما برای دریافت فایل اکسل

# چون ربات روی سرور ابری خارج از ایران اجرا می‌شود، نیازی به پروکسی نیست
bot = Bot(token=TOKEN)
dp = Dispatcher()

class SurveyStates(StatesGroup):
    waiting_for_name = State()
    waiting_for_phone = State()
    waiting_for_activity = State()
    waiting_for_products = State()
    waiting_for_melamine_brands = State()
    waiting_for_mdf_brands = State()
    waiting_for_highgloss_brands = State()
    waiting_for_satisfaction = State()
    waiting_for_feedback_choice = State()
    waiting_for_feedback_text = State()

MELAMINE_BRANDS = [
    "پایتخت", "آذران", "آرتا", "گرندوود", "فومنات", "صنایع",
    "پاک چوب", "نیوپنل", "آرین سینا", "پویا", "وینا", "گنبد",
    "خلخال", "تیسان", "ایزوفام", "ویسپان", "پارس", "چهلستون"
]

MDF_BRANDS = [
    "پایتخت", "آذران", "آرتا", "گرندوود", "فومنات", "صنایع",
    "پاک چوب", "نیوپنل", "آرین سینا", "پویا", "وینا", "گنبد",
    "خلخال", "تیسان", "ایزوفام", "ویسپان", "پارس", "چهلستون", "فرامید"
]

HIGHGLOSS_BRANDS = [
    "دکوپنل", "پانوتک", "S.A.C", "AGT", "اشیک", "یلدیز", "کینگ گلاس", "ایپک"
]

def get_multi_select_kb(brand_list, selected_items):
    builder = InlineKeyboardBuilder()
    for brand in brand_list:
        text = f"✅ {brand}" if brand in selected_items else brand
        builder.button(text=text, callback_data=f"brand_{brand}")
    builder.adjust(3)
    builder.row(types.InlineKeyboardButton(text="ادامه ➡️ تأیید و مرحله بعد", callback_data="next_step"))
    return builder.as_markup()

async def route_to_next_step(state: FSMContext, callback_or_message):
    data = await state.get_data()
    selected_prods = data.get("selected_products", [])
    
    message_to_edit = callback_or_message.message if isinstance(callback_or_message, types.CallbackQuery) else callback_or_message
    
    if "ملامین" in selected_prods:
        await state.update_data(melamine_selected=[])
        kb = get_multi_select_kb(MELAMINE_BRANDS, [])
        await message_to_edit.edit_text(
            "در حال حاضر در بخش ملامین با چه برندهایی فعالیت دارید؟ (چند انتخابی - پس از انتخاب روی ادامه بزنید):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_melamine_brands)
    elif "MDF" in selected_prods:
        await state.update_data(mdf_selected=[])
        kb = get_multi_select_kb(MDF_BRANDS, [])
        await message_to_edit.edit_text(
            "در حال حاضر در بخش MDF از چه برندهایی استفاده می‌کنید؟ (چند انتخابی):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_mdf_brands)
    elif "هایگلاس" in selected_prods:
        await state.update_data(highgloss_selected=[])
        kb = get_multi_select_kb(HIGHGLOSS_BRANDS, [])
        await message_to_edit.edit_text(
            "در حال حاضر در بخش هایگلاس از چه برندهایی استفاده می‌کنید؟ (چند انتخابی):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_highgloss_brands)
    else:
        await ask_satisfaction(callback_or_message, state)

async def ask_satisfaction(callback_or_message, state: FSMContext):
    builder = InlineKeyboardBuilder()
    builder.button(text="عالی ⭐⭐⭐⭐", callback_data="sat_4")
    builder.button(text="خوب ⭐⭐⭐", callback_data="sat_3")
    builder.button(text="متوسط ⭐⭐", callback_data="sat_2")
    builder.button(text="بد ⭐", callback_data="sat_1")
    builder.adjust(2)
    
    text = "میزان رضایت شما از خدمات و محصولات فروشگاه پارس به چه میزان است؟"
    message_to_edit = callback_or_message.message if isinstance(callback_or_message, types.CallbackQuery) else callback_or_message
    await message_to_edit.edit_text(text, reply_markup=builder.as_markup())
    await state.set_state(SurveyStates.waiting_for_satisfaction)

async def finalize_and_send(bot_instance, chat_id, user_id, state: FSMContext, message_obj=None):
    data = await state.get_data()
    
    rows_data = {
        "عنوان فیلد": [
            "نام و نام خانوادگی", "شماره تماس", "نوع فعالیت",
            "محصولات درخواستی", "برندهای ملامین", "برندهای MDF",
            "برندهای هایگلاس", "میزان رضایت", "نظرات و پیشنهادات"
        ],
        "مقدار ثبت‌شده": [
            data.get('name'),
            data.get('phone'),
            data.get('activity'),
            ', '.join(data.get('selected_products', [])) if data.get('selected_products') else "-",
            ', '.join(data.get('melamine_selected', [])) if 'ملامین' in data.get('selected_products', []) else "-",
            ', '.join(data.get('mdf_selected', [])) if 'MDF' in data.get('selected_products', []) else "-",
            ', '.join(data.get('highgloss_selected', [])) if 'هایگلاس' in data.get('selected_products', []) else "-",
            data.get('satisfaction'),
            data.get('feedback', 'ندارد')
        ]
    }
    
    df = pd.DataFrame(rows_data)
    file_name = f"customer_{user_id}.xlsx"
    
    with pd.ExcelWriter(file_name, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Report')
        worksheet = writer.sheets['Report']
        worksheet.sheet_view.rightToLeft = True
        
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(name="Tahoma", size=11, bold=True, color="FFFFFF")
        cell_font = Font(name="Tahoma", size=10)
        
        for col in range(1, 3):
            cell = worksheet.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
        for row in range(2, len(df) + 2):
            for col in range(1, 3):
                cell = worksheet.cell(row=row, column=col)
                cell.font = cell_font
                cell.alignment = Alignment(horizontal="right" if col == 2 else "center", vertical="center")
                
        worksheet.column_dimensions['A'].width = 22
        worksheet.column_dimensions['B'].width = 35

    # 1. ارسال فایل اکسل به آیدی ادمین
    try:
        document = types.FSInputFile(file_name)
        await bot_instance.send_document(
            chat_id=ADMIN_ID,
            document=document,
            caption=f"📊 گزارش سریع مشتری جدید\n👤 نام: {data.get('name')}\n📞 شماره: {data.get('phone')}"
        )
    except Exception as e:
        print(f"Error sending document to admin: {e}")
    
    # 2. ارسال پیام موفقیت و کد تخفیف به مشتری
    name = data.get('name', 'کاربر')
    summary = (
        f"✅ {name} عزیز، ثبت‌نام و اطلاعات شما با موفقیت در سیستم فروشگاه پارس ثبت شد!\n\n"
        f"🎁 کد تخفیف شما برای تمامی محصولات کینگ گلاس: offpersis"
    )
    
    if message_obj:
        if isinstance(message_obj, types.CallbackQuery):
            await message_obj.message.edit_text(summary)
        else:
            await message_obj.answer(summary)
            
    if os.path.exists(file_name):
        os.remove(file_name)
        
    await state.clear()

@dp.message(CommandStart())
async def cmd_start(message: types.Message, state: FSMContext):
    await state.clear()
    await message.answer("لطفاً نام و نام خانوادگی خود را وارد کنید:")
    await state.set_state(SurveyStates.waiting_for_name)

@dp.message(SurveyStates.waiting_for_name)
async def process_name(message: types.Message, state: FSMContext):
    await state.update_data(name=message.text)
    builder = ReplyKeyboardBuilder()
    builder.button(text="اشتراک‌گذاری شماره موبایل 📱", request_contact=True)
    await message.answer(
        "ممنون. لطفاً برای ادامه شماره همراه خود را از طریق دکمه زیر ارسال کنید:",
        reply_markup=builder.as_markup(resize_keyboard=True, one_time_keyboard=True)
    )
    await state.set_state(SurveyStates.waiting_for_phone)

@dp.message(SurveyStates.waiting_for_phone, F.contact)
async def process_phone(message: types.Message, state: FSMContext):
    phone = message.contact.phone_number
    await state.update_data(phone=phone)
    
    builder = InlineKeyboardBuilder()
    builder.button(text="کابینت‌ساز", callback_data="act_cabinet")
    builder.button(text="مهندسین سازه", callback_data="act_structure")
    builder.button(text="مهندس ساختمان", callback_data="act_building")
    builder.button(text="طراحی داخلی", callback_data="act_interior")
    builder.button(text="خانه‌دار", callback_data="act_housewife")
    builder.adjust(2)
    
    await message.answer("نوع فعالیت خود را انتخاب کنید:", reply_markup=types.ReplyKeyboardRemove())
    await message.answer("لطفاً از گزینه‌های زیر انتخاب کنید:", reply_markup=builder.as_markup())
    await state.set_state(SurveyStates.waiting_for_activity)

@dp.callback_query(SurveyStates.waiting_for_activity, F.data.startswith("act_"))
async def process_activity(callback: types.CallbackQuery, state: FSMContext):
    activity_map = {
        "act_cabinet": "کابینت‌ساز",
        "act_structure": "مهندسین سازه",
        "act_building": "مهندس ساختمان",
        "act_interior": "طراحی داخلی",
        "act_housewife": "خانه‌دار"
    }
    selected_act = activity_map.get(callback.data, "نامشخص")
    await state.update_data(activity=selected_act)
    
    if callback.data == "act_housewife":
        await state.update_data(selected_products=[])
        await callback.answer()
        await ask_satisfaction(callback, state)
    else:
        await state.update_data(selected_products=[])
        builder = InlineKeyboardBuilder()
        products = ["هایگلاس", "ملامین", "MDF", "ابزار‌های صنعتی", "تیغه‌های برش"]
        for p in products:
            builder.button(text=f"[ ] {p}", callback_data=f"prod_{p}")
        builder.adjust(2)
        builder.row(types.InlineKeyboardButton(text="تأیید و ادامه ➡️", callback_data="done_products"))
        
        await callback.message.edit_text(
            "متقاضی کدام یک از محصولات فروشگاه پارس می‌باشید؟ (چند انتخابی - پس از انتخاب روی تأیید بزنید):",
            reply_markup=builder.as_markup()
        )
        await callback.answer()
        await state.set_state(SurveyStates.waiting_for_products)

@dp.callback_query(SurveyStates.waiting_for_products, F.data.startswith("prod_"))
async def toggle_product(callback: types.CallbackQuery, state: FSMContext):
    prod_name = callback.data.split("_", 1)[1]
    data = await state.get_data()
    selected = data.get("selected_products", [])
    
    if prod_name in selected:
        selected.remove(prod_name)
    else:
        selected.append(prod_name)
        
    await state.update_data(selected_products=selected)
    
    builder = InlineKeyboardBuilder()
    products = ["هایگلاس", "ملامین", "MDF", "ابزار‌های صنعتی", "تیغه‌های برش"]
    for p in products:
        prefix = "✅ " if p in selected else "[ ] "
        builder.button(text=f"{prefix}{p}", callback_data=f"prod_{p}")
    builder.adjust(2)
    builder.row(types.InlineKeyboardButton(text="تأیید و ادامه ➡️", callback_data="done_products"))
    
    await callback.message.edit_reply_markup(reply_markup=builder.as_markup())
    await callback.answer()

@dp.callback_query(SurveyStates.waiting_for_products, F.data == "done_products")
async def finish_products(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    selected_prods = data.get("selected_products", [])
    
    if not selected_prods:
        await callback.answer("⚠️ کاربر محترم، حتماً باید حداقل یک گزینه را انتخاب کنید!", show_alert=True)
        return
        
    await callback.answer()
    await route_to_next_step(state, callback)

# ملامین
@dp.callback_query(SurveyStates.waiting_for_melamine_brands, F.data.startswith("brand_"))
async def toggle_melamine(callback: types.CallbackQuery, state: FSMContext):
    brand = callback.data.split("_", 1)[1]
    data = await state.get_data()
    selected = data.get("melamine_selected", [])
    if brand in selected:
        selected.remove(brand)
    else:
        selected.append(brand)
    await state.update_data(melamine_selected=selected)
    kb = get_multi_select_kb(MELAMINE_BRANDS, selected)
    await callback.message.edit_reply_markup(reply_markup=kb)
    await callback.answer()

@dp.callback_query(SurveyStates.waiting_for_melamine_brands, F.data == "next_step")
async def finish_melamine(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    selected_brands = data.get("melamine_selected", [])
    
    if not selected_brands:
        await callback.answer("⚠️ کاربر محترم، حتماً باید حداقل یک گزینه را انتخاب کنید!", show_alert=True)
        return
        
    await callback.answer()
    selected_prods = data.get("selected_products", [])
    
    if "MDF" in selected_prods:
        await state.update_data(mdf_selected=[])
        kb = get_multi_select_kb(MDF_BRANDS, [])
        await callback.message.edit_text(
            "در حال حاضر در بخش MDF از چه برندهایی استفاده می‌کنید؟ (چند انتخابی):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_mdf_brands)
    elif "هایگلاس" in selected_prods:
        await state.update_data(highgloss_selected=[])
        kb = get_multi_select_kb(HIGHGLOSS_BRANDS, [])
        await callback.message.edit_text(
            "در حال حاضر در بخش هایگلاس از چه برندهایی استفاده می‌کنید؟ (چند انتخابی):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_highgloss_brands)
    else:
        await ask_satisfaction(callback, state)

# MDF
@dp.callback_query(SurveyStates.waiting_for_mdf_brands, F.data.startswith("brand_"))
async def toggle_mdf(callback: types.CallbackQuery, state: FSMContext):
    brand = callback.data.split("_", 1)[1]
    data = await state.get_data()
    selected = data.get("mdf_selected", [])
    if brand in selected:
        selected.remove(brand)
    else:
        selected.append(brand)
    await state.update_data(mdf_selected=selected)
    kb = get_multi_select_kb(MDF_BRANDS, selected)
    await callback.message.edit_reply_markup(reply_markup=kb)
    await callback.answer()

@dp.callback_query(SurveyStates.waiting_for_mdf_brands, F.data == "next_step")
async def finish_mdf(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    selected_brands = data.get("mdf_selected", [])
    
    if not selected_brands:
        await callback.answer("⚠️ کاربر محترم، حتماً باید حداقل یک گزینه را انتخاب کنید!", show_alert=True)
        return
        
    await callback.answer()
    selected_prods = data.get("selected_products", [])
    
    if "هایگلاس" in selected_prods:
        await state.update_data(highgloss_selected=[])
        kb = get_multi_select_kb(HIGHGLOSS_BRANDS, [])
        await callback.message.edit_text(
            "در حال حاضر در بخش هایگلاس از چه برندهایی استفاده می‌کنید؟ (چند انتخابی):",
            reply_markup=kb
        )
        await state.set_state(SurveyStates.waiting_for_highgloss_brands)
    else:
        await ask_satisfaction(callback, state)

# هایگلاس
@dp.callback_query(SurveyStates.waiting_for_highgloss_brands, F.data.startswith("brand_"))
async def toggle_highgloss(callback: types.CallbackQuery, state: FSMContext):
    brand = callback.data.split("_", 1)[1]
    data = await state.get_data()
    selected = data.get("highgloss_selected", [])
    if brand in selected:
        selected.remove(brand)
    else:
        selected.append(brand)
    await state.update_data(highgloss_selected=selected)
    kb = get_multi_select_kb(HIGHGLOSS_BRANDS, selected)
    await callback.message.edit_reply_markup(reply_markup=kb)
    await callback.answer()

@dp.callback_query(SurveyStates.waiting_for_highgloss_brands, F.data == "next_step")
async def finish_highgloss(callback: types.CallbackQuery, state: FSMContext):
    data = await state.get_data()
    selected_brands = data.get("highgloss_selected", [])
    
    if not selected_brands:
        await callback.answer("⚠️ کاربر محترم، حتماً باید حداقل یک گزینه را انتخاب کنید!", show_alert=True)
        return
        
    await callback.answer()
    await ask_satisfaction(callback, state)

# سنجش میزان رضایت (ستاره‌ها)
@dp.callback_query(SurveyStates.waiting_for_satisfaction, F.data.startswith("sat_"))
async def process_satisfaction(callback: types.CallbackQuery, state: FSMContext):
    sat_map = {"sat_4": "عالی ⭐⭐⭐⭐", "sat_3": "خوب ⭐⭐⭐", "sat_2": "متوسط ⭐⭐", "sat_1": "بد ⭐"}
    await state.update_data(satisfaction=sat_map.get(callback.data, ""))
    
    builder = InlineKeyboardBuilder()
    builder.button(text="پیشنهاد/انتقاد دارم ✍️", callback_data="has_feedback")
    builder.button(text="پیشنهاد/انتقاد ندارم ❌", callback_data="no_feedback")
    builder.adjust(1)
    
    await callback.message.edit_text("آیا پیشنهاد یا انتقادی نسبت به خدمات ما دارید؟", reply_markup=builder.as_markup())
    await callback.answer()
    await state.set_state(SurveyStates.waiting_for_feedback_choice)

# انتخاب کاربر: ندارم
@dp.callback_query(SurveyStates.waiting_for_feedback_choice, F.data == "no_feedback")
async def process_no_feedback(callback: types.CallbackQuery, state: FSMContext):
    await state.update_data(feedback="ندارد")
    await finalize_and_send(bot, callback.message.chat.id, callback.from_user.id, state, callback)

# انتخاب کاربر: دارم
@dp.callback_query(SurveyStates.waiting_for_feedback_choice, F.data == "has_feedback")
async def process_has_feedback(callback: types.CallbackQuery, state: FSMContext):
    await callback.message.edit_text("لطفاً پیشنهاد، انتقاد یا نظر خود را بنویسید:")
    await callback.answer()
    await state.set_state(SurveyStates.waiting_for_feedback_text)

# دریافت متن پیام کاربر
@dp.message(SurveyStates.waiting_for_feedback_text)
async def process_feedback_text_input(message: types.Message, state: FSMContext):
    await state.update_data(feedback=message.text)
    
    builder = InlineKeyboardBuilder()
    builder.button(text="ارسال پیشنهاد/انتقاد 📤", callback_data="submit_final_feedback")
    
    await message.answer("متن شما دریافت شد. برای اتمام کار روی دکمه زیر بزنید:", reply_markup=builder.as_markup())

# دکمه ارسال نهایی پیشنهاد/انتقاد
@dp.callback_query(SurveyStates.waiting_for_feedback_text, F.data == "submit_final_feedback")
async def submit_final_feedback(callback: types.CallbackQuery, state: FSMContext):
    await callback.answer()
    await finalize_and_send(bot, callback.message.chat.id, callback.from_user.id, state, callback)

async def main():
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())